"""Specify shared-capture report behavior without depending on OOXML spelling.

These regression sources cover immutable captures, identical paired outputs,
physical layer mapping, template semantics, and safe output failure paths.
"""

from base64 import b64decode
from collections.abc import Sequence
from copy import copy
from dataclasses import FrozenInstanceError, replace
from decimal import Decimal
import hashlib
from pathlib import Path
import re
import struct
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import zlib

import pytest

# Load the plugin's bundled dependencies without relying on test import order.
from impedance import html_report, service, workbook
from impedance.matching import analyze
from impedance.model import (
    BoardSnapshot,
    Config,
    LayerSettings,
    Section,
    Specification,
    Trace,
    ValidationError,
)

# isort: split
# The plugin adds its bundled openpyxl to the import path above.
from openpyxl import load_workbook
from openpyxl.drawing import image as spreadsheet_image
from openpyxl.reader import drawings as spreadsheet_drawings

_DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"


def _chunk(kind: bytes, payload: bytes) -> bytes:
    """Encode a standard PNG fixture chunk, not a production PNG validator."""
    return (
        struct.pack(">I", len(payload))
        + kind
        + payload
        + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)
    )


def _png(width: int = 12, height: int = 6, color: bytes = b"\x40\x80\xc0") -> bytes:
    """Produce a small RGB fixture without requiring Pillow or KiCad."""
    pixels = (b"\x00" + color * width) * height
    return (
        b"\x89PNG\r\n\x1a\n"
        + _chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + _chunk(b"IDAT", zlib.compress(pixels))
        + _chunk(b"IEND", b"")
    )


def _row(**changes: Any) -> service.ReportRow:
    """Provide an inner-layer 90-ohm USB CPWG with two ground references."""
    values = {
        "specification": "USB CPWG",
        "nets": ("USB_D+", "USB_D-"),
        "net_class": "USB",
        "copper_layers": ("F.Cu", "In1.Cu", "In2.Cu", "B.Cu"),
        "signal_layer": "In1.Cu",
        "reference_layers": ("F.Cu", "In2.Cu"),
        "kind": "differential_coplanar",
        "width_nm": 290576,
        "spacing_nm": 203200,
        "ground_gap_nm": 254000,
        "target_ohms": "90",
        "image": service.CapturedImage.from_bytes(_png()),
    }
    values.update(changes)
    return service.ReportRow(**values)


def _media(path: Path) -> list[bytes]:
    """Read embedded PNGs without requiring particular generated member names."""
    with ZipFile(path) as archive:
        return [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("xl/media/") and name.endswith(".png")
        ]


def _html_images(path: Path) -> list[bytes]:
    """Extract the report's self-contained image payloads in document order."""
    return [
        b64decode(encoded, validate=True)
        for encoded in re.findall(
            r'src="data:image/png;base64,([A-Za-z0-9+/=]+)"',
            path.read_text(encoding="utf-8"),
        )
    ]


def test_captured_image_is_immutable_and_fingerprints_exact_bytes() -> None:
    """A capture's dimensions and identity belong to its owned original PNG."""
    data = _png(30, 10)
    captured = service.CapturedImage.from_bytes(data)
    assert (captured.data, captured.width, captured.height) == (data, 30, 10)
    assert captured.sha256 == hashlib.sha256(data).hexdigest()
    with pytest.raises(FrozenInstanceError):
        setattr(captured, "width", 99)
    with pytest.raises(FrozenInstanceError):
        setattr(_row(image=captured), "width_nm", 1)


@pytest.mark.parametrize("width", [None, True, False])
def test_report_row_requires_an_actual_integer_width(width: Any) -> None:
    """A missing width or Python boolean cannot become a manufacturing number."""
    with pytest.raises(ValidationError, match="actual trace width"):
        _row(width_nm=width)


@pytest.mark.parametrize(
    "data", [b"", b"not a PNG", b"\x89PNG\r\n\x1a\n", _png(0, 1), _png(4097, 1)]
)
def test_captured_image_rejects_missing_or_unbounded_headers(data: bytes) -> None:
    """Bound locally produced captures without reimplementing a PNG decoder."""
    with pytest.raises(ValidationError):
        service.CapturedImage.from_bytes(data)


def test_capture_byte_limit_applies_to_bytes_and_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Neither capture entry point accepts an oversized renderer result."""
    path = tmp_path / "capture.png"
    path.write_bytes(_png())
    monkeypatch.setattr(service, "MAX_IMAGE_BYTES", 16)
    with pytest.raises(ValidationError):
        service.CapturedImage.from_bytes(path.read_bytes())
    with pytest.raises(ValidationError):
        service.CapturedImage.load(path)


@pytest.mark.parametrize("remove_source", [False, True])
def test_both_reports_use_owned_image_after_source_changes(
    tmp_path: Path, remove_source: bool
) -> None:
    """Changing or deleting a temporary PNG cannot change either report."""
    path = tmp_path / "capture.png"
    original = _png(24, 8)
    path.write_bytes(original)
    captured = service.CapturedImage.load(path)
    if remove_source:
        path.unlink()
    else:
        path.write_bytes(_png(8, 24, b"\xff\x00\x00"))
    rows = (_row(image=captured),)
    xlsx = workbook.write_workbook(rows, tmp_path / "report.xlsx")
    html = html_report.write_html_report(rows, tmp_path / "report.html")
    assert _media(xlsx) == _html_images(html) == [original]
    assert hashlib.sha256(_media(xlsx)[0]).hexdigest() == captured.sha256


@pytest.mark.parametrize(
    ("signal", "references", "physical", "physical_references"),
    [
        ("F.Cu", ("In1.Cu",), "L1", "L2"),
        ("In1.Cu", ("F.Cu", "In2.Cu"), "L2", "L1, L3"),
        ("B.Cu", ("In2.Cu",), "L4", "L3"),
    ],
)
def test_reports_map_physical_layers_and_keep_actual_width_and_both_gaps(
    tmp_path: Path,
    signal: str,
    references: tuple[str, ...],
    physical: str,
    physical_references: str,
) -> None:
    """Both reference planes and pair/ground gaps survive the canonical row."""
    row = _row(signal_layer=signal, reference_layers=references)
    xlsx = workbook.write_workbook([row], tmp_path / "report.xlsx")
    html = html_report.write_html_report([row], tmp_path / "report.html")
    sheet = load_workbook(xlsx).active
    assert sheet["A2"].value == physical
    assert sheet["B2"].value == physical_references
    assert sheet["C2"].value == "Differential Pair (Coplanar)"
    assert Decimal(str(sheet["D2"].value)) == Decimal("11.44")
    assert sheet["F2"].value == "Pair: 8\nGround: 10"
    assert sheet["G2"].value == 90
    markup = html.read_text(encoding="utf-8")
    assert f"{signal} ({physical})" in markup
    for canonical, physical_reference in zip(references, row.physical_reference_layers):
        assert f"{canonical} ({physical_reference})" in markup
    assert "0.290576 mm" in markup and "11.44 mil" in markup
    assert "290576 nm" not in markup


@pytest.mark.parametrize(
    ("kind", "label", "spacing"),
    [
        ("single_ended", "Single Ended (Non coplanar)", None),
        ("single_ended_coplanar", "Single Ended (Coplanar)", 10),
        ("differential", "Differential Pair (Non coplanar)", 8),
        (
            "differential_coplanar",
            "Differential Pair (Coplanar)",
            "Pair: 8\nGround: 10",
        ),
    ],
)
def test_workbook_uses_vendor_types_and_only_applicable_spacing(
    tmp_path: Path, kind: str, label: str, spacing: Any
) -> None:
    """Keep a single-ended CPWG ground gap distinct from differential spacing."""
    differential = kind.startswith("differential")
    row = _row(
        kind=kind,
        target_ohms="90" if differential else "50",
        nets=("USB_D+", "USB_D-") if differential else ("RF",),
        spacing_nm=203200 if differential else None,
        ground_gap_nm=254000 if kind.endswith("_coplanar") else None,
    )
    output = workbook.write_workbook([row], tmp_path / "report.xlsx")
    sheet = load_workbook(output).active
    assert sheet["C2"].value == label
    assert sheet["F2"].value == spacing


def test_template_header_style_comments_and_original_are_preserved(
    tmp_path: Path,
) -> None:
    """Preserve vendor meaning and formatting, not ZIP/XML byte serialization."""
    source_bytes = workbook.TEMPLATE_PATH.read_bytes()
    original = load_workbook(workbook.TEMPLATE_PATH).active
    output = workbook.write_workbook([_row()] * 5, tmp_path / "report.xlsx")
    generated = load_workbook(output).active
    for before, after in zip(original[1], generated[1]):
        assert after.value == before.value
        assert after.number_format == before.number_format
        for attribute in ("font", "fill", "border", "alignment", "protection"):
            assert copy(getattr(after, attribute)) == copy(getattr(before, attribute))
        if before.comment is not None:
            assert after.comment is not None
            assert (after.comment.text, after.comment.author) == (
                before.comment.text,
                before.comment.author,
            )
    # The real vendor note belongs to C2, not the header row. It lists the
    # allowed impedance types and must survive replacing the example rows.
    before_comment = original["C2"].comment
    after_comment = generated["C2"].comment
    assert before_comment is not None and after_comment is not None
    assert (after_comment.text, after_comment.author) == (
        before_comment.text,
        before_comment.author,
    )
    assert generated["C3"].comment is None
    assert generated.max_row == 6
    assert {str(area) for area in generated.merged_cells.ranges} == {
        f"D{index}:E{index}" for index in range(1, 7)
    }
    assert "$A$1:$H$6" in str(generated.print_area)
    assert generated.print_title_rows == "$1:$1"
    assert generated.page_setup.fitToWidth == 1
    assert generated.page_setup.fitToHeight == 0
    assert len(_media(output)) == 5
    assert workbook.TEMPLATE_PATH.read_bytes() == source_bytes


@pytest.mark.parametrize(("width", "height"), [(800, 420), (420, 800), (2, 1000)])
def test_workbook_image_keeps_aspect_ratio_and_fits_its_row(
    tmp_path: Path, width: int, height: int
) -> None:
    """Landscape, portrait, and narrow captures must not stretch or overlap rows."""
    row = _row(image=service.CapturedImage.from_bytes(_png(width, height)))
    output = workbook.write_workbook([row], tmp_path / "report.xlsx")
    sheet = load_workbook(output).active
    with ZipFile(output) as archive:
        drawing_name = next(
            name
            for name in archive.namelist()
            if re.fullmatch(r"xl/drawings/drawing\d+\.xml", name)
        )
        drawing = ET.fromstring(archive.read(drawing_name))  # noqa: S314
    anchor = drawing.find(f"{{{_DRAWING}}}oneCellAnchor")
    assert anchor is not None
    extent = anchor.find(f"{{{_DRAWING}}}ext")
    assert extent is not None
    cx, cy = int(extent.attrib["cx"]), int(extent.attrib["cy"])
    assert cx > 0 and cy > 0
    assert abs(cx * height - cy * width) <= max(width, height)
    assert cy < sheet.row_dimensions[2].height * 12700
    assert anchor.findtext(f"{{{_DRAWING}}}from/{{{_DRAWING}}}col") == "7"
    assert anchor.findtext(f"{{{_DRAWING}}}from/{{{_DRAWING}}}row") == "1"


def test_png_embedding_does_not_require_pillow(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The platform-neutral PNG adapter must not enter Pillow's image loader."""
    monkeypatch.setattr(spreadsheet_image, "PILImage", None)
    monkeypatch.setattr(spreadsheet_drawings, "PILImage", None)
    row = _row()
    output = workbook.write_workbook([row], tmp_path / "report.xlsx")
    assert _media(output) == [row.image.data]


def test_display_metadata_is_escaped_and_cells_cannot_be_formulas(
    tmp_path: Path,
) -> None:
    """Names stay literal in HTML; manufacturing cells contain mapped/numeric data."""
    row = _row(
        specification='<script>alert("USB")</script>',
        nets=('=HYPERLINK("https://invalid")', "USB<&D-"),
        net_class="<USB>&RF",
    )
    xlsx = workbook.write_workbook([row], tmp_path / "report.xlsx")
    html = html_report.write_html_report(
        [row], tmp_path / "report.html", board_name="<board>&USB"
    )
    sheet = load_workbook(xlsx).active
    assert all(cell.data_type != "f" for cells in sheet for cell in cells)
    markup = html.read_text(encoding="utf-8")
    assert "<script>" not in markup
    assert "&lt;script&gt;" in markup and "USB&lt;&amp;D-" in markup
    assert "&lt;board&gt;&amp;USB" in markup


@pytest.mark.parametrize("current", [False, True])
def test_html_preserves_review_times_without_approving_historical_records(
    tmp_path: Path, current: bool
) -> None:
    """An export displays supplied review currentness without minting approval."""
    row = _row(
        last_viewed_at_utc="2026-09-08T13:14:15.000000Z",
        last_viewed_current=current,
        layer_approved_at_utc="2026-09-08T13:15:16.000000Z",
        layer_approval_current=current,
    )
    before = replace(row)
    output = html_report.write_html_report([row], tmp_path / "report.html")
    markup = output.read_text(encoding="utf-8")
    assert "2026-09-08 13:14:15 UTC" in markup
    assert "2026-09-08 13:15:16 UTC" in markup
    assert ("Current capture" if current else "Previous capture") in markup
    assert (
        "Current layer settings and captures"
        if current
        else "Previous layer settings or captures"
    ) in markup
    assert row == before


@pytest.mark.parametrize(
    "writer", [workbook.write_workbook, html_report.write_html_report]
)
@pytest.mark.parametrize("limit", ["empty", "rows", "images"])
def test_report_resource_limits_preserve_previous_outputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, writer: Any, limit: str
) -> None:
    """Both writers enforce the same complete-report limits before publication."""
    row = _row()
    rows = [row, row]
    if limit == "empty":
        rows = []
    elif limit == "rows":
        monkeypatch.setattr(service, "MAX_ROWS", 1)
    else:
        monkeypatch.setattr(service, "MAX_TOTAL_IMAGE_BYTES", len(row.image.data))
    output = tmp_path / "previous.output"
    output.write_bytes(b"previous report")
    with pytest.raises(ValidationError):
        writer(rows, output)
    assert output.read_bytes() == b"previous report"
    assert set(tmp_path.iterdir()) == {output}


@pytest.mark.parametrize("alias", [False, True])
def test_workbook_refuses_to_replace_its_template(tmp_path: Path, alias: bool) -> None:
    """Protect both direct and symlink template destinations using a copied fixture."""
    template = tmp_path / "template.xlsx"
    source = workbook.TEMPLATE_PATH.read_bytes()
    template.write_bytes(source)
    destination = template
    if alias:
        destination = tmp_path / "alias.xlsx"
        destination.symlink_to(template)
    with pytest.raises(ValidationError, match="overwrite.*template"):
        workbook.write_workbook([_row()], destination, template=template)
    assert template.read_bytes() == source


@pytest.mark.parametrize(
    "writer", [workbook.write_workbook, html_report.write_html_report]
)
def test_report_failure_keeps_previous_output_and_cleans_temporary_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, writer: Any
) -> None:
    """Atomic publication failure must not destroy the previous report."""
    output = tmp_path / "report.output"
    output.write_bytes(b"previous report")

    def fail_replace(_source: Any, _destination: Any) -> None:
        """Simulate a filesystem error after the complete report was prepared."""
        raise OSError("publication unavailable")

    monkeypatch.setattr(workbook.os, "replace", fail_replace)
    with pytest.raises(OSError, match="publication unavailable"):
        writer([_row()], output)
    assert output.read_bytes() == b"previous report"
    assert set(tmp_path.iterdir()) == {output}


@pytest.mark.parametrize(
    ("target", "width", "spacing", "ground_gap", "message"),
    [
        ("1e20", 203200, 203200, 254000, "Target impedance"),
        ("90", 10**12 + 1, 203200, 254000, "actual trace width"),
        ("90", 203200, 10**12 + 1, 254000, "Report gaps"),
        ("90", 203200, 203200, 10**12 + 1, "Report gaps"),
    ],
)
def test_vendor_bounds_fail_preflight_before_rendering_or_writing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    target: str,
    width: int,
    spacing: int,
    ground_gap: int,
    message: str,
) -> None:
    """Domain-valid values outside report bounds must fail before fabrication work."""
    snapshot = BoardSnapshot(
        ("F.Cu", "B.Cu"),
        (
            Trace("plus", "F.Cu", "USB_D+", width, ((0, 0), (1000000, 0))),
            Trace(
                "minus",
                "F.Cu",
                "USB_D-",
                width,
                ((0, 406400), (1000000, 406400)),
            ),
        ),
        net_classes=("USB",),
        net_class_context_digest="usb-class-constraints",
        net_class_memberships=(("USB_D+", ("USB",)), ("USB_D-", ("USB",))),
        differential_pairs=(("USB_D+", "USB_D-"),),
    )
    config = Config(
        enabled=True,
        specifications=(
            Specification(
                "usb",
                "USB CPWG",
                target,
                "differential_coplanar",
                "USB",
                (LayerSettings("F.Cu", ("B.Cu",), spacing, ground_gap),),
            ),
        ),
    )
    analysis = analyze(config, snapshot)
    assert len(analysis.sections) == 1
    config = replace(
        config,
        reviewed_digest=analysis.digest,
        included_section_ids=(analysis.sections[0].section_id,),
    )

    def unexpected_work(*_args: Any, **_kwargs: Any) -> Any:
        """Fail immediately if invalid report metadata reaches a renderer or writer."""
        pytest.fail("Preflight must reject vendor bounds before output work.")

    monkeypatch.setattr(html_report, "write_html_report", unexpected_work)
    with pytest.raises(ValidationError, match=message):
        service.prepare(config, snapshot, layer_count=2)
    plan = service.ExportPlan(config, snapshot, analysis, analysis.sections)
    with pytest.raises(ValidationError, match=message):
        service.export_reports(
            plan,
            object(),
            object(),
            tmp_path,
            renderer_factory=unexpected_work,
            writer=unexpected_work,
        )
    assert not list(tmp_path.iterdir())


def test_workbook_write_failure_closes_owned_stream_before_cleanup(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An exception retaining openpyxl's writer must not retain an open output file."""
    from openpyxl.workbook import Workbook

    destination = tmp_path / "report.xlsx"
    destination.write_bytes(b"previous report")
    streams: list[Any] = []

    def fail_save(_workbook: Workbook, stream: Any) -> None:
        """Retain the output stream as a failed library traceback would."""
        streams.append(stream)
        stream.write(b"unfinished workbook")
        raise OSError("spreadsheet write failed")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError, match="spreadsheet write failed") as failure:
        workbook.write_workbook([_row()], destination)
    assert failure.value.__traceback__ is not None
    assert len(streams) == 1 and streams[0].closed
    assert destination.read_bytes() == b"previous report"
    assert set(tmp_path.iterdir()) == {destination}


def test_export_pipeline_loads_each_capture_once_and_shares_canonical_rows(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Real preflight/export shares immutable rows and releases the renderer."""
    snapshot = BoardSnapshot(
        ("F.Cu", "In1.Cu", "In2.Cu", "B.Cu"),
        (
            Trace("short", "F.Cu", "RF_SHORT", 180000, ((0, 0), (1000000, 0))),
            Trace(
                "long",
                "F.Cu",
                "RF_LONG",
                180000,
                ((0, 10000000), (120000000, 10000000)),
            ),
        ),
        net_classes=("RF",),
        net_class_context_digest="rf-class-constraints",
        net_class_memberships=(("RF_SHORT", ("RF",)), ("RF_LONG", ("RF",))),
    )
    config = Config(
        enabled=True,
        specifications=(
            Specification(
                "rf",
                "RF routes",
                "50",
                "single_ended",
                "RF",
                (LayerSettings("F.Cu", ("In1.Cu",)),),
            ),
        ),
    )
    analysis = analyze(config, snapshot)
    config = replace(
        config,
        reviewed_digest=analysis.digest,
        included_section_ids=tuple(section.section_id for section in analysis.sections),
    )
    plan = service.prepare(config, snapshot, layer_count=4)
    assert plan is not None and len(plan.sections) == 2
    paths: list[Path] = []
    loads: list[Path] = []
    seen: list[Sequence[service.ReportRow]] = []
    original_load = service.CapturedImage.load
    original_workbook = workbook.write_workbook
    original_html = html_report.write_html_report

    class Renderer:
        """Produce deterministic distinct captures and track native-style cleanup."""

        closed = False

        def render(
            self, section: Section, destination: Path, width_px: int, height_px: int
        ) -> Path:
            """Generate a PNG for each complete section in export order."""
            color = (
                b"\x10\x20\x30"
                if section.net_names == ("RF_LONG",)
                else b"\x90\x80\x70"
            )
            destination.write_bytes(_png(20, 10, color))
            paths.append(destination)
            return destination

        def close(self) -> None:
            """Record release of the renderer-owned detached board."""
            self.closed = True

    renderer = Renderer()

    def load_once(_cls: type, path: Path) -> service.CapturedImage:
        """Count actual file capture loads without replacing their implementation."""
        loads.append(path)
        return original_load(path)

    def write_xlsx(rows: Sequence[service.ReportRow], destination: Path) -> Path:
        """Remove temporary sources after embedding the canonical captures."""
        seen.append(rows)
        result = original_workbook(rows, destination)
        for path in paths:
            path.unlink()
        return result

    def write_html(
        rows: Sequence[service.ReportRow], destination: Path, **options: Any
    ) -> Path:
        """Use the real companion writer after every source PNG has disappeared."""
        seen.append(rows)
        return original_html(rows, destination, **options)

    monkeypatch.setattr(service.CapturedImage, "load", classmethod(load_once))
    monkeypatch.setattr(html_report, "write_html_report", write_html)
    artifacts = service.export_reports(
        plan,
        object(),
        object(),
        tmp_path,
        renderer_factory=lambda _board, _pcbnew: renderer,
        writer=write_xlsx,
    )
    assert renderer.closed
    assert loads == paths
    assert len(seen) == 2 and seen[0] is seen[1]
    assert all(isinstance(row, service.ReportRow) for row in seen[0])
    expected = [row.image.data for row in seen[0]]
    assert _media(artifacts.workbook) == _html_images(artifacts.html_report) == expected


@pytest.mark.parametrize(
    "change", ["none", "source", "image", "settings", "omitted_sibling"]
)
def test_report_rows_keep_audit_times_but_recompute_exact_review_currentness(
    change: str,
) -> None:
    """Historical times cannot approve changed pixels, source, intent or row sets."""
    from impedance.review_tracking import (
        ImageView,
        LayerApproval,
        ReviewTracking,
        capture_fingerprint,
        layer_fingerprint,
        source_fingerprint,
    )

    snapshot = BoardSnapshot(
        ("F.Cu", "In1.Cu", "In2.Cu", "B.Cu"),
        (
            Trace("first", "F.Cu", "RF_A", 180000, ((0, 0), (1000000, 0))),
            Trace(
                "second", "F.Cu", "RF_B", 180000, ((0, 5000000), (120000000, 5000000))
            ),
            Trace("back", "B.Cu", "RF_C", 180000, ((0, 8000000), (1000000, 8000000))),
        ),
        context_digest="original-native-source",
        net_classes=("RF",),
        net_class_memberships=tuple(
            (name, ("RF",)) for name in ("RF_A", "RF_B", "RF_C")
        ),
        net_class_context_digest="rf-constraints",
    )
    spec = Specification(
        "rf",
        "RF routes",
        "50",
        "single_ended",
        "RF",
        (LayerSettings("F.Cu", ("In1.Cu",)), LayerSettings("B.Cu", ("In2.Cu",))),
    )
    config = Config(enabled=True, specifications=(spec,))
    analysis = analyze(config, snapshot)
    source = source_fingerprint(snapshot)
    images = tuple(
        service.CapturedImage.from_bytes(_png(color=bytes((index, 50, 100))))
        for index in range(len(analysis.sections))
    )
    captures = tuple(
        (section.section_id, capture_fingerprint(source, spec, section, image.sha256))
        for section, image in zip(analysis.sections, images)
    )
    viewed_at = "2026-09-08T13:14:15.000000Z"
    approved_at = "2026-09-08T13:15:16.000000Z"
    tracking = ReviewTracking(
        tuple(
            ImageView(
                spec.spec_id, section.layer, section.section_id, digest, viewed_at
            )
            for section, (_, digest) in zip(analysis.sections, captures)
        ),
        tuple(
            LayerApproval(
                spec.spec_id,
                layer,
                layer_fingerprint(source, spec, layer, analysis.sections),
                approved_at,
                tuple(
                    capture
                    for section, capture in zip(analysis.sections, captures)
                    if section.layer == layer
                ),
            )
            for layer in ("F.Cu", "B.Cu")
        ),
    )
    config = replace(
        config,
        reviewed_digest=analysis.digest,
        included_section_ids=tuple(section.section_id for section in analysis.sections),
        review_tracking=tracking,
    )
    plan = service.prepare(config, snapshot)
    assert plan is not None
    before = config.to_dict()
    if change == "source":
        plan = replace(
            plan, snapshot=replace(snapshot, context_digest="changed-zone-fill")
        )
    elif change == "image":
        images = (
            service.CapturedImage.from_bytes(_png(color=b"\xff\xff\x00")),
            *images[1:],
        )
    elif change == "settings":
        spec = replace(
            spec,
            layer_settings=(LayerSettings("F.Cu", ("In2.Cu",)), spec.layer_settings[1]),
        )
        plan = replace(plan, config=replace(config, specifications=(spec,)))
    elif change == "omitted_sibling":
        plan = replace(plan, sections=(analysis.sections[0], analysis.sections[2]))
        images = (images[0], images[2])
    rows = service.report_rows(plan, images)
    for index, row in enumerate(rows):
        assert row.last_viewed_at_utc == viewed_at
        assert row.layer_approved_at_utc == approved_at
        assert row.last_viewed_current is (
            change != "source" and not (change == "image" and index == 0)
        )
        assert row.layer_approval_current is (
            change == "none" or (row.signal_layer == "B.Cu" and change != "source")
        )
    assert config.to_dict() == before
