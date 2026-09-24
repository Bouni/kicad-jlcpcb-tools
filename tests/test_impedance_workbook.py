"""Verify template fidelity, image integrity, and atomic impedance workbooks."""

from copy import copy
from dataclasses import replace
import hashlib
from pathlib import Path
import posixpath
import struct
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import zlib

import pytest

from impedance import service, workbook
from impedance.model import ValidationError

# isort: split
# The plugin adds its bundled openpyxl to the import path above.
from openpyxl import load_workbook

_NS = {
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "x": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def _chunk(kind: bytes, payload: bytes) -> bytes:
    """Encode a PNG chunk with a valid checksum."""
    return (
        struct.pack(">I", len(payload))
        + kind
        + payload
        + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)
    )


def _png(
    width: int = 8, height: int = 4, pixel_data: bytes = b"", interlace: int = 0
) -> bytes:
    """Make a small self-contained RGB PNG using only the standard library."""
    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, interlace)
    pixels = pixel_data or (b"\x00" + b"\x70\x80\x90" * width) * height
    return (
        b"\x89PNG\r\n\x1a\n"
        + _chunk(b"IHDR", header)
        + _chunk(b"IDAT", zlib.compress(pixels))
        + _chunk(b"IEND", b"")
    )


def _xml(data: bytes) -> ET.Element:
    """Read test-generated XML and the hash-verified bundled fixture."""
    return ET.fromstring(data)  # noqa: S314 -- Only known fixture/generated XML.


def _parts(path: Path) -> dict[str, bytes]:
    """Read a workbook's decompressed members for byte-preservation assertions."""
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) == len(set(archive.namelist()))
        return {name: archive.read(name) for name in archive.namelist()}


@pytest.fixture
def row() -> service.ReportRow:
    """Provide one immutable capture with canonical physical layer mapping."""
    return service.ReportRow(
        "RF",
        ("RF",),
        "RF",
        ("F.Cu", "In1.Cu", "In2.Cu", "B.Cu"),
        "F.Cu",
        ("In1.Cu",),
        "single_ended",
        254000,
        None,
        "50",
        service.CapturedImage.from_bytes(_png()),
    )


def test_template_is_exact_supplied_workbook() -> None:
    """Pin the vendor resource, not generated XML serialization."""
    assert (
        hashlib.sha256(workbook.TEMPLATE_PATH.read_bytes()).hexdigest()
        == workbook.TEMPLATE_SHA256
    )
    assert workbook.validate_template() is None


def test_wrapped_row_styles_preserve_vendor_formatting(
    row: service.ReportRow, tmp_path: Path
) -> None:
    """Manufacturing fields wrap without losing vendor formatting."""
    report_row = replace(
        row,
        signal_layer="In1.Cu",
        reference_layers=("F.Cu", "In2.Cu"),
        kind="differential_coplanar",
        spacing_nm=203200,
        ground_gap_nm=254000,
    )
    output = workbook.write_workbook([report_row], tmp_path / "report.xlsx")
    original, generated = load_workbook(workbook.TEMPLATE_PATH), load_workbook(output)
    try:
        for column in ("A", "B", "C", "F"):
            before, after = (
                original.active[f"{column}2"],
                generated.active[f"{column}2"],
            )
            assert after.alignment.wrap_text is True
            for attribute in (
                "horizontal",
                "vertical",
                "text_rotation",
                "shrink_to_fit",
                "indent",
            ):
                assert getattr(after.alignment, attribute) == getattr(
                    before.alignment, attribute
                )
            for attribute in ("font", "fill", "border", "number_format"):
                assert copy(getattr(after, attribute)) == copy(
                    getattr(before, attribute)
                )
        assert generated.active["F2"].value == "Pair: 8\nGround: 10"
        assert generated.active["B2"].value == "L1, L3"
    finally:
        original.close()
        generated.close()


@pytest.mark.parametrize("count", [1, 3, 4, 37, service.MAX_ROWS])
def test_dynamic_rows_images_merges_and_print_area(
    row: service.ReportRow, tmp_path: Path, count: int
) -> None:
    """Each captured row retains an in-row image, merge and printable area."""
    output = workbook.write_workbook([row] * count, tmp_path / "report.xlsx")
    parts = _parts(output)
    sheet = _xml(parts["xl/worksheets/sheet1.xml"])
    rows = sheet.findall("s:sheetData/s:row", _NS)
    assert len(rows) == count + 1
    assert all(float(item.get("ht")) == 255 for item in rows[1:])
    assert {
        item.get("ref") for item in sheet.findall("s:mergeCells/s:mergeCell", _NS)
    } == {f"D{index}:E{index}" for index in range(1, count + 2)}
    anchors = _xml(parts["xl/drawings/drawing1.xml"]).findall("x:oneCellAnchor", _NS)
    relationships = _xml(parts["xl/drawings/_rels/drawing1.xml.rels"])
    targets = {item.get("Id"): item.get("Target") for item in relationships}
    assert len(anchors) == len(targets) == count
    embedded = set()
    for index, anchor in enumerate(anchors, start=1):
        assert anchor.findtext("x:from/x:row", namespaces=_NS) == str(index)
        assert anchor.findtext("x:from/x:col", namespaces=_NS) == "7"
        extent = anchor.find("x:ext", _NS)
        assert 0 < int(extent.get("cx")) <= 600 * 9525
        assert 0 < int(extent.get("cy")) + 57150 < 255 * 12700
        relation = anchor.find("x:pic/x:blipFill/a:blip", _NS).get(
            f"{{{_NS['r']}}}embed"
        )
        target = targets[relation]
        member = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("xl/drawings", target))
        )
        assert parts[member] == row.image.data
        embedded.add(member)
    assert len(embedded) == count
    assert embedded == {name for name in parts if name.startswith("xl/media/")}
    generated = load_workbook(output)
    try:
        sheet = generated.active
        assert "$A$1:$H$" + str(count + 1) in str(sheet.print_area)
        assert sheet.print_title_rows == "$1:$1"
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == 0
    finally:
        generated.close()


@pytest.mark.parametrize(
    "fields",
    [
        {"signal_layer": ""},
        {"signal_layer": "unknown"},
        {"reference_layers": ()},
        {"reference_layers": ("",)},
        {"reference_layers": "In1.Cu"},
        {"reference_layers": ("F.Cu",)},
        {"reference_layers": ("In1.Cu", "In1.Cu")},
        {"copper_layers": ("F.Cu", "F.Cu", "In1.Cu")},
        {"kind": "unknown"},
        {"width_nm": 0},
        {"width_nm": -1},
        {"width_nm": 10**12 + 1},
        {"width_nm": 25400.0},
        {"width_nm": True},
        {"spacing_nm": 0},
        {"spacing_nm": -1},
        {"spacing_nm": 0.1},
        {"ground_gap_nm": 0},
        {"ground_gap_nm": -1},
        {"ground_gap_nm": 0.1},
        {"target_ohms": ""},
        {"target_ohms": "NaN"},
        {"target_ohms": "Infinity"},
        {"target_ohms": "-10"},
        {"target_ohms": "0"},
        {"target_ohms": "1e999999"},
        {"target_ohms": "1e-999999"},
        {"target_ohms": "=50"},
        {"target_ohms": "1" * 65},
        {"kind": "differential", "spacing_nm": None},
        {"kind": "single_ended_coplanar", "ground_gap_nm": None},
        {"kind": "differential_coplanar", "spacing_nm": 254000},
    ],
)
def test_invalid_row_leaves_previous_workbook_intact(
    row: service.ReportRow, tmp_path: Path, fields: dict[str, Any]
) -> None:
    """Invalid metadata fails at immutable construction before publication."""
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"previous valid workbook")
    with pytest.raises((ValidationError, TypeError)):
        workbook.write_workbook([replace(row, **fields)], output)
    assert output.read_bytes() == b"previous valid workbook"
    assert set(tmp_path.iterdir()) == {output}


@pytest.mark.parametrize("count", [0, service.MAX_ROWS + 1])
def test_empty_or_excessive_sections_rejected(
    row: service.ReportRow, tmp_path: Path, count: int
) -> None:
    """Reject missing reviewed rows and bound complete report resources."""
    with pytest.raises(ValidationError, match="captured rows"):
        workbook.write_workbook([row] * count, tmp_path / "report.xlsx")


@pytest.mark.parametrize(
    "payload", [b"not a PNG", _png()[:32], _png(4097, 1), _png(1, 4097), _png(0, 1)]
)
def test_invalid_capture_container_preserves_previous_output(
    row: service.ReportRow, tmp_path: Path, payload: bytes
) -> None:
    """The native decoder owns pixels; the shared boundary checks container bounds."""
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"previous workbook")
    with pytest.raises(ValidationError):
        capture = service.CapturedImage.from_bytes(payload)
        workbook.write_workbook([replace(row, image=capture)], output)
    assert output.read_bytes() == b"previous workbook"


def test_missing_capture_has_actionable_error(tmp_path: Path) -> None:
    """A missing renderer result fails before a report row can be created."""
    with pytest.raises(ValidationError, match="Cannot read impedance capture"):
        service.CapturedImage.load(tmp_path / "missing.png")


@pytest.mark.parametrize(
    "template_data",
    [b"not a zip", b"PK\x03\x04", workbook.TEMPLATE_PATH.read_bytes() + b"changed"],
)
def test_unknown_or_corrupt_template_rejected(
    row: service.ReportRow, tmp_path: Path, template_data: bytes
) -> None:
    """Modified vendor templates cannot publish an output."""
    template = tmp_path / "template.xlsx"
    template.write_bytes(template_data)
    with pytest.raises(ValidationError, match="Unsupported JLCPCB impedance template"):
        workbook.write_workbook([row], tmp_path / "report.xlsx", template)


def test_missing_template(row: service.ReportRow, tmp_path: Path) -> None:
    """Expose a missing packaged resource with an actionable error."""
    with pytest.raises(ValidationError, match="Cannot read JLCPCB impedance template"):
        workbook.write_workbook(
            [row], tmp_path / "report.xlsx", tmp_path / "missing.xlsx"
        )
