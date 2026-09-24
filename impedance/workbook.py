"""Populate JLCPCB's template with openpyxl and the shared captured report rows."""

from collections.abc import Sequence
from copy import copy
from decimal import Decimal
import hashlib
from io import BytesIO
import os
from pathlib import Path
import tempfile
from typing import Any, Optional

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU

from .model import ValidationError
from .service import KINDS, CapturedImage, ReportRow, validate_report_rows

TEMPLATE_SHA256 = "22d94b15293435f103868cb62cdbee3c32db6d7d635344baafa332d630b04bd4"
TEMPLATE_PATH = Path(__file__).parent / "resources" / "Required_impedance_control.xlsx"
OUTPUT_FILENAME = "Required_impedance_control.xlsx"
WorkbookError = ValidationError


class _CapturedPng(Image):
    """Give openpyxl an already-rendered PNG without Pillow or format conversion.

    openpyxl 3.1.5's Image contract is width, height, format, anchor, and _data();
    inherited path/_id handling still belongs to its ordinary drawing writer.
    """

    def __init__(self, capture: CapturedImage) -> None:
        self.width, self.height = capture.width, capture.height
        self.format = "png"
        self._capture = capture

    def _data(self) -> bytes:
        """Embed the exact bytes displayed in the preview and HTML companion."""
        return self._capture.data


def _template_bytes(template: Path) -> bytes:
    """Check the bundled vendor artifact, not serialized output internals."""
    try:
        with template.open("rb") as source:
            data = source.read(2 * 1024 * 1024 + 1)
    except OSError as error:
        raise WorkbookError(
            f"Cannot read JLCPCB impedance template: {template}."
        ) from error
    if hashlib.sha256(data).hexdigest() != TEMPLATE_SHA256:
        raise WorkbookError(
            "Unsupported JLCPCB impedance template; the bundled resource is missing or changed."
        )
    return data


def validate_template(template: Optional[Path] = None) -> None:
    """Fail preflight when the shipped workbook resource cannot be used."""
    _template_bytes(TEMPLATE_PATH if template is None else Path(template))


def _cells(row: ReportRow) -> tuple[Any, ...]:
    """Map one canonical report row to the vendor's eight fixed columns."""
    width = Decimal(row.width_nm) / Decimal(25400)
    spacing = (
        None if row.spacing_nm is None else Decimal(row.spacing_nm) / Decimal(25400)
    )
    ground = (
        None
        if row.ground_gap_nm is None
        else Decimal(row.ground_gap_nm) / Decimal(25400)
    )
    if row.kind == "differential_coplanar":
        spacing = f"Pair: {spacing:f}\nGround: {ground:f}"
    elif row.kind == "single_ended_coplanar":
        spacing = ground
    elif row.kind == "single_ended":
        spacing = None
    return (
        row.physical_signal_layer,
        ", ".join(row.physical_reference_layers),
        KINDS[row.kind],
        width,
        None,
        spacing,
        Decimal(row.target_ohms),
        None,
    )


def write_workbook(
    rows: Sequence[ReportRow], destination: Path, template: Optional[Path] = None
) -> Path:
    """Preserve the template's visible layout and atomically replace report output.

    General XLSX serialization is owned by openpyxl. Untouched ZIP entries need
    not be byte-identical: headers, styles, comments, columns and print layout
    are preserved semantically, and vendor example rows/images are replaced.
    """
    validate_report_rows(rows)
    destination = Path(destination)
    source = TEMPLATE_PATH if template is None else Path(template)
    if destination.resolve() == source.resolve() or (
        destination.exists()
        and source.exists()
        and os.path.samefile(destination, source)
    ):
        raise WorkbookError(
            "The impedance workbook destination must not overwrite its template."
        )
    workbook = load_workbook(BytesIO(_template_bytes(source)), keep_links=False)
    sheet = workbook.active
    styles = [copy(sheet.cell(2, column)._style) for column in range(1, 9)]
    comments = [copy(sheet.cell(2, column).comment) for column in range(1, 9)]
    for merged in tuple(sheet.merged_cells.ranges):
        if merged.min_row > 1:
            sheet.unmerge_cells(str(merged))
    sheet.delete_rows(2, sheet.max_row)
    for index in tuple(sheet.row_dimensions):
        if index > 1:
            del sheet.row_dimensions[index]
    # Clear vendor examples whether or not Pillow happened to load them.
    sheet._images.clear()
    for index, row in enumerate(rows, start=2):
        for column, value in enumerate(_cells(row), start=1):
            cell = sheet.cell(index, column, value)
            cell._style = copy(styles[column - 1])
            if index == 2:
                cell.comment = comments[column - 1]
            if isinstance(value, str):
                cell.data_type = "s"  # User text is never interpreted as a formula.
            if column in (1, 2, 3, 6):
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
        sheet.merge_cells(start_row=index, start_column=4, end_row=index, end_column=5)
        sheet.row_dimensions[index].height = 255
        image = _CapturedPng(row.image)
        scale = min(600 / image.width, 315 / image.height)
        image.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=7, row=index - 1, colOff=pixels_to_EMU(6), rowOff=pixels_to_EMU(6)
            ),
            ext=XDRPositiveSize2D(
                cx=round(image.width * scale * 9525),
                cy=round(image.height * scale * 9525),
            ),
        )
        sheet.add_image(image)
    sheet.print_area = f"A1:H{len(rows) + 1}"
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    temporary: Optional[Path] = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{destination.name}.",
            suffix=".tmp",
            dir=destination.parent,
            delete=False,
        ) as output:
            temporary = Path(output.name)
            # Own the stream even when openpyxl fails before closing its ZIP.
            # This also permits cleanup on platforms that cannot unlink open files.
            workbook.save(output)
            output.flush()
            os.fsync(output.fileno())
        os.replace(temporary, destination)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return destination
